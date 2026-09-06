import requests
import openpyxl
import aiohttp

from bgu_helper.logger import logger

url = "https://docs.google.com/spreadsheets/d/1Bi5nCoVbFyV7i-0ELHBzgozbzlAdzBfP/export?format=xlsx"


# print(sheet.cell(21, 1).value)

async def get_weather():
    async with aiohttp.ClientSession() as session:
        async with session.get(
            "https://api.open-meteo.com/v1/forecast",
            params={
                "latitude": 53.9,
                "longitude": 27.5667,
                "current": "temperature_2m"
            }
        ) as response:

            response.raise_for_status()

            data = await response.json()

            return data["current"]["temperature_2m"]


def get_schedule() -> dict[str, dict[str, list[str | None]]]:
    try:
        response = requests.get(url)
        response.raise_for_status()

        with open("schedule.xlsx", "wb") as f:
            f.write(response.content)

        book = openpyxl.load_workbook("schedule.xlsx")

        sheet = book.active

        day = None
        schedule = {}

        for row in range(21, sheet.max_row + 1):

            if sheet.cell(row, 1).value is not None:
                day = sheet.cell(row, 1).value
                schedule[day] = {}
                # print(day)

            time = sheet.cell(row, 2).value

            if time is not None:

                if sheet.cell(row, 16).value is not None:
                    schedule[day][time] = []
                    schedule[day][time].append(sheet.cell(row, 16).value)
                    schedule[day][time].append(sheet.cell(row+1, 16).value)
                    schedule[day][time].append(sheet.cell(row+3, 16).value)
                elif sheet.cell(row, 12).value is not None and sheet.cell(row+3, 12).value is None:
                    schedule[day][time] = []
                    schedule[day][time].append(sheet.cell(row, 12).value)
                    schedule[day][time].append(sheet.cell(row+2, 12).value)
                    schedule[day][time].append(sheet.cell(row+3, 16).value)
        return schedule
    except Exception as e:
        logger.error(f"Ошибка получения расписания: {e}")
        return {}


def load_schedule(scheduale):
    rows_for_schedule = ""
    try:
        for day in scheduale:
            rows_for_schedule += f"""
                <tr>
                    <td colspan="2"><b>{day.upper()}</b></td>
                </tr>
                """
            for time in scheduale[day]:
                rows_for_schedule += f"""
                <tr>
                    <td>{time}</td>
                    <td>
                            <p>{scheduale[day][time][0]}</p>
                            <p>{scheduale[day][time][1] if scheduale[day][time][1] is not None else "нет данных"} | {scheduale[day][time][2] if scheduale[day][time][2] is not None else "нет данных"}</p>
                    </td>
                </tr>
                """
        return rows_for_schedule
    except Exception as e:
        logger.error(f"Ошибка получения расписания: {e}")
        return '<tr><td colspan="2">Ошибка получения расписания</td></tr>'


def print_schedule(schedule):
    for day in schedule:
        print(day)
        for time in schedule[day]:
            print('\t', time)
            for lesson in schedule[day][time]:
                print('\t\t', lesson)


# print_schedule(schedule)
